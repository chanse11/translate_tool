"""PDF → Word：按文本块重建段落，避免误识别表格。"""

from __future__ import annotations

import sys
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _lines_to_rows(lines: list) -> list[str]:
    """同一 y 坐标的行横向拼接；列间距较大时用 Tab 分隔。"""
    from collections import defaultdict

    rows: dict[float, list[tuple[float, float, str]]] = defaultdict(list)
    for line in lines:
        y = round(line["bbox"][1], 1)
        text = "".join(span["text"] for span in line.get("spans", []))
        if text:
            rows[y].append((line["bbox"][0], line["bbox"][2], text))

    row_texts: list[str] = []
    for y in sorted(rows.keys()):
        parts = sorted(rows[y], key=lambda item: item[0])
        chunks: list[str] = [parts[0][2]]
        for (_, prev_x1, _), (curr_x0, _, curr_text) in zip(parts, parts[1:]):
            gap = curr_x0 - prev_x1
            chunks.append("\t" + curr_text if gap > 24 else curr_text)
        row_texts.append("".join(chunks))
    return row_texts


def _block_text(block: dict) -> str:
    rows = _lines_to_rows(block.get("lines", []))
    return "".join(rows).strip()


def _block_bottom(block: dict) -> float:
    return block["bbox"][3]


def _block_top(block: dict) -> float:
    return block["bbox"][1]


def _looks_like_section_start(text: str) -> bool:
    stripped = text.lstrip()
    if stripped.startswith("第") and "条" in stripped[:6]:
        return True
    if stripped in {"房屋租赁合同"}:
        return True
    return False


def _merge_blocks(blocks: list[dict]) -> list[dict]:
    """合并同一自然段内被拆开的文本块。"""
    if not blocks:
        return []

    merged: list[dict] = [blocks[0]]
    for block in blocks[1:]:
        prev = merged[-1]
        gap = _block_top(block) - _block_bottom(prev)
        prev_text = _block_text(prev)
        curr_text = _block_text(block)

        same_paragraph = (
            gap <= 18
            and not _looks_like_section_start(curr_text)
            and not prev_text.endswith("：")
        )
        if same_paragraph:
            prev["lines"].extend(block.get("lines", []))
            prev["bbox"] = (
                min(prev["bbox"][0], block["bbox"][0]),
                min(prev["bbox"][1], block["bbox"][1]),
                max(prev["bbox"][2], block["bbox"][2]),
                max(prev["bbox"][3], block["bbox"][3]),
            )
        else:
            merged.append(block)
    return merged


def _block_font_size(block: dict) -> float | None:
    sizes: list[float] = []
    for line in block.get("lines", []):
        for span in line.get("spans", []):
            if span.get("text", "").strip():
                sizes.append(span.get("size", 0))
    return max(sizes) if sizes else None


def _is_bold(block: dict) -> bool:
    for line in block.get("lines", []):
        for span in line.get("spans", []):
            flags = span.get("flags", 0)
            if flags & 2**4:  # bold
                return True
    return False


def _guess_alignment(block: dict, page_width: float) -> WD_ALIGN_PARAGRAPH:
    x0, _, x1, _ = block["bbox"]
    center = (x0 + x1) / 2
    page_center = page_width / 2
    if abs(center - page_center) < page_width * 0.08 and (x1 - x0) < page_width * 0.5:
        return WD_ALIGN_PARAGRAPH.CENTER
    if x0 > page_width * 0.12:
        return WD_ALIGN_PARAGRAPH.LEFT
    return WD_ALIGN_PARAGRAPH.JUSTIFY


def _extract_text_blocks(page: fitz.Page) -> list[dict]:
    blocks = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        text = _block_text(block)
        if not text:
            continue
        blocks.append(block)
    blocks.sort(key=lambda b: (round(b["bbox"][1], 1), b["bbox"][0]))
    return blocks


def convert_pdf_to_docx(pdf_path: Path, docx_path: Path) -> None:
    pdf = fitz.open(pdf_path)
    doc = Document()

    # 默认正文样式
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)

    for page_index, page in enumerate(pdf):
        if page_index > 0:
            doc.add_page_break()

        page_width = page.rect.width
        blocks = _merge_blocks(_extract_text_blocks(page))

        for block in blocks:
            text = _block_text(block)
            para = doc.add_paragraph()
            para.alignment = _guess_alignment(block, page_width)

            run = para.add_run(text)
            size = _block_font_size(block)
            if size:
                run.font.size = Pt(round(size, 1))
            if _is_bold(block):
                run.bold = True

            # 标题通常更大
            if size and size >= 16:
                run.bold = True

    pdf.close()
    doc.save(docx_path)


def main() -> None:
    if len(sys.argv) < 2:
        print("用法: python pdf_to_docx_paragraphs.py <input.pdf> [output.docx]")
        sys.exit(1)

    pdf_path = Path(sys.argv[1]).resolve()
    if len(sys.argv) >= 3:
        docx_path = Path(sys.argv[2]).resolve()
    else:
        docx_path = pdf_path.with_suffix(".docx").with_stem(pdf_path.stem + "-段落版")

    convert_pdf_to_docx(pdf_path, docx_path)
    print(f"已生成: {docx_path}")


if __name__ == "__main__":
    main()
